import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import schedule
from schedule import build_xlsx
from tests.helpers import make_person, temporary_workbook_template


class WorkbookTests(unittest.TestCase):
    def test_build_xlsx_writes_source_sheet_and_values(self):
        people = [
            make_person("Alex", [{"start": "11:00", "end": "11:30"}]),
            make_person("Blake", [{"start": "14:00", "end": "14:30"}]),
        ]
        room_data = {
            "11:00": {"off": "Officer A", "b": "3", "s": "3"},
            "14:00": {"off": "Officer B", "b": "2", "s": "1"},
        }

        with temporary_workbook_template():
            buffer = build_xlsx(people, room_data)

        result = load_workbook(buffer)
        self.addCleanup(result.close)
        self.assertIn("Source Data", result.sheetnames)
        self.assertIn("AM", result.sheetnames)
        self.assertIn("PM", result.sheetnames)

        source = result["Source Data"]
        self.assertEqual(source["A1"].value, "Alex, 11:00, 11:30")
        self.assertEqual(source["A2"].value, "Blake, 14:00, 14:30")
        self.assertEqual(source["A4"].value, "11:00, Officer A, 3, 3")

    def test_build_xlsx_escapes_formula_like_names_in_source_sheet(self):
        people = [
            {"name": "=cmd|' /C calc'!A0", "ranges": [{"start": "11:00", "end": "11:30"}]},
        ]

        with temporary_workbook_template():
            buffer = build_xlsx(people, {})

        exported = load_workbook(buffer)
        self.addCleanup(exported.close)
        self.assertEqual(exported["AM"]["C7"].data_type, "s")
        self.assertTrue(exported["AM"]["C7"].value.startswith("'=cmd"))
        self.assertEqual(exported["Source Data"]["A1"].value, "'=cmd|' /C calc'!A0, 11:00, 11:30")

    def test_build_xlsx_raises_when_template_is_missing(self):
        with patch.object(schedule, "TEMPLATE_PATH", "/definitely/missing/template.xlsx"):
            with self.assertRaises(FileNotFoundError):
                build_xlsx(
                    [make_person("Alex", [{"start": "11:00", "end": "11:30"}])],
                    {},
                )

    def test_schedule_cells_filter_sort_shade_and_include_room_headers(self):
        people = [
            make_person("Zoe", [{"start": "11:30", "end": "12:00"}]),
            make_person("Blake", [{"start": "14:00", "end": "14:30"}]),
            make_person("Alex", [{"start": "11:00", "end": "11:30"}]),
        ]
        with temporary_workbook_template():
            buffer = build_xlsx(people, {"14:00": {"off": "Officer B", "b": "2", "s": "1"}})
        result = load_workbook(buffer)
        self.addCleanup(result.close)
        am, pm = result["AM"], result["PM"]
        self.assertEqual(
            [am[f"C{row}"].value for row in (6, 7, 8, 9)], ["Off.", "Alex", "Zoe", None]
        )
        self.assertEqual([pm[f"C{row}"].value for row in (6, 7, 8)], ["Off.", "Blake", None])
        self.assertEqual(am["D7"].fill.patternType, None)
        self.assertEqual(am["P7"].fill.fgColor.rgb, "00EDEDED")
        self.assertEqual(am["D8"].fill.fgColor.rgb, "00EDEDED")
        self.assertEqual(am["P8"].fill.patternType, None)
        self.assertEqual([am[cell].value for cell in ("D26", "P26", "AB26")], [2, 2, 1])
        self.assertEqual(
            [pm[cell].value for cell in ("G4", "F3", "L3", "H3", "N3")],
            ["Officer B", "2", "1", "0B", "1S"],
        )

    def test_shipped_template_exports_schedule_and_preserves_layout(self):
        result = load_workbook(
            build_xlsx(
                [
                    make_person("Alex", [{"start": "11:00", "end": "11:30"}]),
                    make_person("Blake", [{"start": "14:00", "end": "14:30"}]),
                ],
                {"14:00": {"off": "Officer B"}},
            )
        )
        self.addCleanup(result.close)
        self.assertEqual(result["AM1"]["C7"].value, "Alex")
        self.assertEqual(result["PM1"]["C7"].value, "Blake")
        self.assertEqual(result["PM1"]["G4"].value, "Officer B")
        self.assertEqual(result["AM1"]["D26"].value, 2)
        self.assertTrue(result["AM1"].merged_cells.ranges)
        self.assertIn("Source Data", result.sheetnames)


if __name__ == "__main__":
    unittest.main()
